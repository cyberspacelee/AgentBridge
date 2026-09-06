import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from office import execute, read_office


class OfficeChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="agentbridge-office-")
        self.root = Path(self.temp.name).resolve()

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_formatted_runs_and_literal_table_export(self):
        source = self.root / "input.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Start 20").bold = True
        paragraph.add_run("25 End")
        for table_number in range(2):
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "=1+1"
            table.cell(1, 1).text = str(table_number)
        doc.save(source)
        result = execute(self.root, {"action": "replace", "path": "input.docx", "output": "updated.docx", "replacements": {"2025": "2026"}})
        self.assertEqual(result["replacements"], 1)
        changed = Document(self.root / "updated.docx")
        self.assertEqual(changed.paragraphs[0].text, "Start 2026 End")
        self.assertTrue(changed.paragraphs[0].runs[0].bold)
        with zipfile.ZipFile(source) as a, zipfile.ZipFile(self.root / "updated.docx") as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != "word/document.xml":
                    self.assertEqual(a.read(name), b.read(name))
        execute(self.root, {"action": "tables_to_xlsx", "path": "input.docx", "output": "tables.xlsx"})
        book = load_workbook(self.root / "tables.xlsx")
        self.assertEqual(book.sheetnames, ["Table_1", "Table_2"])
        self.assertEqual(book.worksheets[0]["A2"].value, "=1+1")
        self.assertEqual(book.worksheets[0]["A2"].data_type, "s")
        book.close()

    def test_pptx_creation_reorder_and_reopen(self):
        execute(self.root, {"action": "create_pptx", "output": "input.pptx", "slides": [{"title": title, "bullets": ["Content"]} for title in ["A", "B", "C"]]})
        execute(self.root, {"action": "replace", "path": "input.pptx", "output": "reordered.pptx", "replacements": {"Content": "Updated"}, "order": [3, 1, 2]})
        presentation = Presentation(self.root / "reordered.pptx")
        self.assertEqual([slide.shapes.title.text for slide in presentation.slides], ["C", "A", "B"])
        self.assertIn("Updated", read_office(self.root / "reordered.pptx")["slides"][0]["text"])
        result = execute(self.root, {"action": "validate", "path": "reordered.pptx"})
        self.assertEqual(result["validation"], "passed")
        with self.assertRaises(ValueError):
            execute(self.root, {"action": "replace", "path": "input.pptx", "output": "bad.pptx", "order": [1, 1, 3]})
        self.assertFalse((self.root / "bad.pptx").exists())

    def test_decimal_aggregation_and_markdown(self):
        (self.root / "sales.csv").write_text("Region,Revenue\nEast,0.1\nEast,0.2\nWest,5\n", encoding="utf-8")
        result = execute(self.root, {"action": "analyse", "path": "sales.csv", "output": "report.md", "groupBy": ["Region"], "measure": "Revenue"})
        self.assertEqual(result["rows"], [["East", "0.3"], ["West", "5"]])
        self.assertIn("| East | 0.3 |", (self.root / "report.md").read_text())

    def test_path_boundaries_no_overwrite_and_scoped_delete(self):
        (self.root / "keep.txt").write_text("keep")
        (self.root / "delete.txt").write_text("delete")
        self.assertEqual(execute(self.root, {"action": "find", "pattern": "*.txt"})["paths"], ["delete.txt", "keep.txt"])
        result = execute(self.root, {"action": "delete", "paths": ["delete.txt"]})
        self.assertEqual(result["deleted"], ["delete.txt"])
        self.assertTrue((self.root / "keep.txt").exists())
        with self.assertRaises(ValueError):
            execute(self.root, {"action": "read", "path": "../outside.txt"})
        with self.assertRaises(ValueError):
            execute(self.root, {"action": "analyse", "path": "keep.txt", "output": "keep.txt"})
        self.assertEqual((self.root / "keep.txt").read_text(), "keep")


if __name__ == "__main__":
    unittest.main()
