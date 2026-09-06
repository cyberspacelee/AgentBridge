"""Shared, bounded JSON tool entry point. Paths are confined to --directory."""
import argparse
import csv
import fnmatch
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook
from pptx import Presentation

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
XML = etree.XMLParser(resolve_entities=False, no_network=True, remove_blank_text=False)
MAX_BYTES = 100 * 1024 * 1024


def bounded_path(root, value, output=False):
    if not isinstance(value, str) or not value:
        raise ValueError("A non-empty path is required")
    candidate = (root / value).resolve()
    if not candidate.is_relative_to(root) or candidate == root:
        raise ValueError("Path is outside the working directory")
    if output:
        if not candidate.parent.is_dir():
            raise ValueError("Output parent directory does not exist")
        if candidate.exists():
            raise ValueError("Output already exists; choose a new filename")
    elif not candidate.is_file() or candidate.stat().st_size > MAX_BYTES:
        raise ValueError("Input is missing, not a file, or exceeds 100 MiB")
    return candidate


def save_atomic(target, writer):
    descriptor, name = tempfile.mkstemp(prefix=".agentbridge-", suffix=target.suffix, dir=target.parent)
    os.close(descriptor)
    try:
        writer(name)
        with open(name, "r+b") as source:
            os.fsync(source.fileno())
        # Same-directory hard link publishes atomically and never overwrites a racing writer.
        os.link(name, target)
    finally:
        Path(name).unlink(missing_ok=True)
    with target.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    return {"path": str(target), "sizeBytes": target.stat().st_size, "sha256": digest}


def package(source):
    archive = zipfile.ZipFile(source)
    entries = archive.infolist()
    if len(entries) > 10000 or sum(i.file_size for i in entries) > MAX_BYTES or len({i.filename for i in entries}) != len(entries):
        archive.close()
        raise ValueError("Office package exceeds limits or contains duplicate entries")
    return archive


def replace_paragraph(paragraph, text_tag, replacements):
    count = 0
    nodes = list(paragraph.iter(text_tag))
    for old, new in replacements.items():
        text = "".join(node.text or "" for node in nodes)
        matches = list(re.finditer(re.escape(old), text))
        # Apply from right to left so offsets remain valid across formatted runs.
        for match in reversed(matches):
            offset = 0
            for node in nodes:
                value = node.text or ""
                end = offset + len(value)
                left, right = max(offset, match.start()), min(end, match.end())
                if left < right:
                    node.text = value[:left - offset] + (new if offset <= match.start() < end else "") + value[right - offset:]
                    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                offset = end
            count += 1
    return count


def edit_package(source, target, replacements, order=None):
    if not isinstance(replacements, dict) or any(not isinstance(k, str) or not k or not isinstance(v, str) for k, v in replacements.items()):
        raise ValueError("replacements must map non-empty text to replacement text")
    changed, replacements_count = [], 0
    with package(source) as archive:
        updates = {}
        for item in archive.infolist():
            word = item.filename.startswith("word/") and item.filename.endswith(".xml")
            slide = item.filename.startswith("ppt/slides/slide") and item.filename.endswith(".xml")
            if not (word or slide):
                continue
            document = etree.fromstring(archive.read(item), XML)
            prefix = "w" if word else "a"
            count = sum(replace_paragraph(p, f"{{{NS[prefix]}}}t", replacements) for p in document.iter(f"{{{NS[prefix]}}}p"))
            if count:
                updates[item.filename] = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone=True)
                changed.append(item.filename)
                replacements_count += count
        if order is not None:
            document = etree.fromstring(archive.read("ppt/presentation.xml"), XML)
            slides = document.find("p:sldIdLst", NS)
            if slides is None or not isinstance(order, list) or any(type(n) is not int for n in order) or sorted(order) != list(range(1, len(slides) + 1)):
                raise ValueError("order must contain every slide number exactly once (1-based)")
            original = list(slides)
            slides[:] = [original[n - 1] for n in order]
            updates["ppt/presentation.xml"] = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone=True)
            changed.append("ppt/presentation.xml")
        def writer(filename):
            with zipfile.ZipFile(filename, "w") as result:
                for item in archive.infolist():
                    result.writestr(item, updates.get(item.filename, archive.read(item)))
        result = save_atomic(target, writer)
    return {**result, "replacements": replacements_count, "changedParts": changed}


def read_office(source):
    if source.suffix.lower() == ".docx":
        with package(source):
            doc = Document(source)
        return {"paragraphs": [p.text for p in doc.paragraphs], "tables": [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]}
    if source.suffix.lower() == ".pptx":
        with package(source):
            presentation = Presentation(source)
        return {"slides": [{"number": index + 1, "text": [shape.text for shape in slide.shapes if shape.has_text_frame], "tables": [[[cell.text for cell in row.cells] for row in shape.table.rows] for shape in slide.shapes if shape.has_table]} for index, slide in enumerate(presentation.slides)]}
    if source.suffix.lower() in (".xlsx", ".csv"):
        return {"sheets": read_sheets(source)}
    raise ValueError("Supported inputs: docx, pptx, xlsx, csv")


def read_sheets(source):
    cells = 0
    def collect(iterator):
        nonlocal cells
        rows = []
        for row in iterator:
            cells += len(row)
            if cells > 200000:
                raise ValueError("Workbook exceeds 200,000-cell tool limit")
            rows.append(list(row))
        return rows
    if source.suffix.lower() == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            rows = collect(csv.reader(stream))
        result = {source.stem: rows}
    else:
        with package(source):
            workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            result = {sheet.title: collect(sheet.iter_rows(values_only=True)) for sheet in workbook}
        finally:
            workbook.close()
    return result


def export_tables(source, target):
    tables = read_office(source)["tables"]
    if not tables:
        raise ValueError("Document has no top-level tables")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, table in enumerate(tables):
        sheet = workbook.create_sheet(f"Table_{index + 1}")
        for row in table:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
    return {**save_atomic(target, workbook.save), "sheets": workbook.sheetnames}


def create_presentation(target, slides):
    if not isinstance(slides, list) or not 1 <= len(slides) <= 5:
        raise ValueError("slides must contain 1 to 5 pages")
    presentation = Presentation()
    for spec in slides:
        if not isinstance(spec, dict) or set(spec) - {"title", "bullets"} or not isinstance(spec.get("title"), str) or not isinstance(spec.get("bullets", []), list):
            raise ValueError("Each slide needs a title and optional bullets")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        body = slide.placeholders[1].text_frame
        for index, bullet in enumerate(spec.get("bullets", [])):
            if not isinstance(bullet, str):
                raise ValueError("Bullets must be text")
            (body.paragraphs[0] if index == 0 else body.add_paragraph()).text = bullet
    return {**save_atomic(target, presentation.save), "slides": len(slides)}


def analyse(source, target, request):
    sheets = read_sheets(source)
    sheet = request.get("sheet") or next(iter(sheets))
    if sheet not in sheets or not sheets[sheet]:
        raise ValueError("Selected sheet is missing or empty")
    rows = sheets[sheet]
    header = [str(v) for v in rows[0]]
    groups = request.get("groupBy", [])
    measure = request.get("measure")
    operation = request.get("aggregate", "sum")
    if not isinstance(groups, list) or any(g not in header for g in groups) or len(set(header)) != len(header):
        raise ValueError("groupBy must reference unique header columns")
    if operation not in ("sum", "mean", "count", "min", "max") or (operation != "count" and measure not in header):
        raise ValueError("aggregate or measure is invalid")
    totals = defaultdict(list)
    for row in rows[1:]:
        if not any(value is not None and value != "" for value in row):
            continue
        padded = row + [None] * max(0, len(header) - len(row))
        key = tuple(str(padded[header.index(group)]) for group in groups)
        try:
            value = Decimal(1) if operation == "count" else Decimal(str(padded[header.index(measure)]))
        except InvalidOperation as error:
            raise ValueError("Measure contains missing, formula, or non-numeric data; clean explicitly before aggregation") from error
        if not value.is_finite():
            raise ValueError("Measure contains a non-finite number")
        totals[key].append(value)
    result = []
    for key, values in sorted(totals.items()):
        value = {"sum": lambda: sum(values), "mean": lambda: sum(values) / len(values), "count": lambda: len(values), "min": lambda: min(values), "max": lambda: max(values)}[operation]()
        result.append([*key, str(value)])
    escape = lambda value: str(value).replace("|", "\\|").replace("\n", " ")
    columns = [*groups, f"{operation}({measure or '*'})"]
    report = f"# Data analysis\n\nSource: {source.name}\nSheet: {sheet}\n\n| " + " | ".join(map(escape, columns)) + " |\n| " + " | ".join("---" for _ in columns) + " |\n"
    report += "\n".join("| " + " | ".join(map(escape, row)) + " |" for row in result) + "\n"
    return {**save_atomic(target, lambda name: Path(name).write_text(report, encoding="utf-8")), "columns": columns, "rows": result}


def validate(source):
    suffix = source.suffix.lower()
    if suffix in (".docx", ".pptx", ".xlsx"):
        with package(source) as archive:
            corrupt = archive.testzip()
            if corrupt:
                raise ValueError("Package CRC validation failed")
            for name in archive.namelist():
                if name.endswith((".xml", ".rels")):
                    etree.fromstring(archive.read(name), XML)
        read_office(source)
    else:
        source.read_text(encoding="utf-8-sig")
    with source.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    return {"path": str(source), "sha256": digest, "validation": "passed", "scope": "package-and-reopen"}


def execute(root, request):
    if not isinstance(request, dict) or not isinstance(request.get("action"), str):
        raise ValueError("A JSON object with action is required")
    action = request["action"]
    allowed = {"read": {"path"}, "replace": {"path", "output", "replacements", "order"}, "tables_to_xlsx": {"path", "output"}, "create_pptx": {"output", "slides"}, "analyse": {"path", "output", "sheet", "groupBy", "measure", "aggregate"}, "validate": {"path"}, "find": {"pattern"}, "delete": {"paths"}, "launch": {"application"}, "search": {"query", "count"}}
    if action not in allowed or set(request) - allowed[action] - {"action"}:
        raise ValueError("Unknown action or unsupported fields")
    source = bounded_path(root, request["path"]) if "path" in request else None
    target = bounded_path(root, request["output"], output=True) if "output" in request else None
    expected = {"replace": source.suffix if source else None, "tables_to_xlsx": ".xlsx", "create_pptx": ".pptx", "analyse": ".md"}
    if action in expected and (target is None or target.suffix.lower() != expected[action]):
        raise ValueError("Output extension does not match the action")
    if action == "read":
        return read_office(source)
    if action == "replace":
        if source.suffix.lower() not in (".docx", ".pptx"):
            raise ValueError("replace supports docx and pptx")
        return edit_package(source, target, request.get("replacements", {}), request.get("order"))
    if action == "tables_to_xlsx":
        return export_tables(source, target)
    if action == "create_pptx":
        return create_presentation(target, request.get("slides"))
    if action == "analyse":
        return analyse(source, target, request)
    if action == "validate":
        return validate(source)
    if action == "find":
        result, visited = [], 0
        pattern = request.get("pattern", "*")
        if not isinstance(pattern, str):
            raise ValueError("pattern must be text")
        for directory, dirs, names in os.walk(root, followlinks=False):
            dirs[:] = [d for d in dirs if d not in (".git", ".venv", "node_modules") and not (Path(directory) / d).is_symlink()]
            visited += len(names) + len(dirs)
            if visited > 10000:
                raise ValueError("Directory search exceeds 10,000 entries")
            for name in names:
                candidate = Path(directory) / name
                if not candidate.is_symlink() and fnmatch.fnmatch(name, pattern):
                    result.append(str(candidate.relative_to(root)))
        return {"paths": sorted(result)}
    if action == "delete":
        values = request.get("paths")
        if not isinstance(values, list) or not 1 <= len(values) <= 1000:
            raise ValueError("paths must contain 1 to 1000 literal file paths")
        paths = list(dict.fromkeys(bounded_path(root, value) for value in values))
        deleted = []
        for candidate in paths:
            try:
                candidate.unlink()
                deleted.append(str(candidate.relative_to(root)))
            except OSError:
                return {"deleted": deleted, "error": "Deletion partially failed; inspect before retrying"}
        return {"deleted": deleted}
    if action == "launch":
        if sys.platform != "win32" or request.get("application") != "outlook":
            raise ValueError("Only Outlook launch on Windows is supported")
        child = subprocess.Popen(["OUTLOOK.EXE"], cwd=root, shell=False)
        return {"pid": child.pid, "status": "started", "application": "outlook"}
    if action == "search":
        token = os.environ.get("BRAVE_SEARCH_API_KEY")
        if not token:
            raise ValueError("BRAVE_SEARCH_API_KEY is not configured")
        query, count = request.get("query"), request.get("count", 5)
        if not isinstance(query, str) or not 1 <= len(query) <= 400 or type(count) is not int or not 1 <= count <= 20:
            raise ValueError("Invalid search query or result count")
        url = "https://api.search.brave.com/res/v1/web/search?" + urllib.parse.urlencode({"q": query, "count": count})
        request = urllib.request.Request(url, headers={"X-Subscription-Token": token, "Accept": "application/json"})
        with urllib.request.urlopen(request, timeout=30) as response:
            data = json.loads(response.read(2 * 1024 * 1024))
        return {"query": query, "results": [{"title": r.get("title"), "url": r.get("url"), "description": r.get("description")} for r in data.get("web", {}).get("results", [])]}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", required=True)
    parser.add_argument("--request", help="JSON request file; omit to read stdin")
    args = parser.parse_args()
    try:
        root = Path(args.directory).resolve(strict=True)
        if not root.is_dir():
            raise ValueError("Working directory must exist")
        raw = bounded_path(root, args.request).read_text(encoding="utf-8") if args.request else sys.stdin.read(1048577)
        if len(raw.encode("utf-8")) > 1048576:
            raise ValueError("Request exceeds 1 MiB")
        result = execute(root, json.loads(raw))
        print(json.dumps({"ok": True, "result": result}, ensure_ascii=True, default=str))
    except (ValueError, KeyError, TypeError) as error:
        print(json.dumps({"ok": False, "code": "VALIDATION_ERROR", "message": str(error)}, ensure_ascii=True))
        return 1
    except Exception as error:
        print(json.dumps({"ok": False, "code": "TOOL_ERROR", "message": type(error).__name__}))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
